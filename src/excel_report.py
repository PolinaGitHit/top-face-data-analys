from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


def _excel_value(v):
    if v is None or (isinstance(v, float) and (np.isnan(v) or np.isinf(v))):
        return None
    return v


def _write_primer_table(ws, first_col: str, data_rows, top: int = 1):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = top
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
    ws.cell(r, 1, first_col).alignment = align
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.cell(r, 2, "Пролонгации в первый месяц").alignment = align
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    ws.cell(r, 5, "Пролонгации через месяц").alignment = align

    sub = ("к пролонгации ", "пролонгировано", "Коэффициент")
    for j, t in zip(range(2, 5), sub):
        ws.cell(r + 1, j, t).alignment = align
    for j, t in zip(range(5, 8), sub):
        ws.cell(r + 1, j, t).alignment = align

    for k, row in enumerate(data_rows):
        for j, v in enumerate(row, start=1):
            c = ws.cell(r + 2 + k, j, _excel_value(v))
            if j > 1 and isinstance(v, (int, float, np.floating)) and not isinstance(v, (bool, np.bool_)):
                c.alignment = align


def _write_coef_block(ws, block_title, col_labels, managers, mat: pd.DataFrame, top: int):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = top
    last_col = 1 + len(col_labels)

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(r, 1, block_title).alignment = align
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 2, end_column=1)
    ws.cell(r + 1, 1, "Менеджер").alignment = align
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=last_col)
    ws.cell(r + 1, 2, "Месяц").alignment = align

    for j, lab in enumerate(col_labels, start=2):
        ws.cell(r + 2, j, lab).alignment = align
    for i, mgr in enumerate(managers):
        ws.cell(r + 3 + i, 1, mgr)
        for j, lab in enumerate(col_labels, start=2):
            v = mat.loc[mgr, lab]
            ws.cell(r + 3 + i, j, _excel_value(v if pd.notna(v) else None)).alignment = align


def _write_rows(ws, headers, rows):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header).alignment = align
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row_idx, col_idx, _excel_value(value)).alignment = align


def _autosize_columns(ws, max_width: int = 42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), max_width)


def _write_summary_sheet(wb: Workbook, result):
    ws = wb.create_sheet("Итоги")
    k1 = result.year_dept_k1
    k2 = result.year_dept_k2
    rows = [
        ("K1 за 2023", k1["den"], k1["num"], round(k1["num"] / k1["den"], 4) if k1["den"] else None),
        ("K2 за 2023", k2["den"], k2["num"], round(k2["num"] / k2["den"], 4) if k2["den"] else None),
    ]
    _write_rows(
        ws,
        ("Метрика", "К пролонгации", "Пролонгировано", "Коэффициент"),
        rows,
    )

    start = len(rows) + 4
    ws.cell(start, 1, "Контрольные показатели")
    control_rows = [
        ("Уникальных проектов в prolongations", len(result.proj)),
        ("Исключено по стоп/end или отсутствию данных", len(result.exclude_stop_set)),
        ("Менеджеров в отчёте", len(result.managers)),
        ("Строк в финансовой агрегации", len(result.fin_long)),
    ]
    for idx, row in enumerate(control_rows, start=start + 1):
        ws.cell(idx, 1, row[0])
        ws.cell(idx, 2, row[1])
    _autosize_columns(ws)


def _write_manager_ranking_sheet(wb: Workbook, result):
    rows = []
    for mgr, k1_den, k1_num, k1_coef, k2_den, k2_num, k2_coef in result.mgr_rows:
        rows.append((mgr, k1_den, k1_num, k1_coef, k2_den, k2_num, k2_coef))

    rows.sort(key=lambda row: (row[3] is None or pd.isna(row[3]), -(row[3] or 0), row[0]))
    ws = wb.create_sheet("Рейтинг менеджеров")
    _write_rows(
        ws,
        (
            "Менеджер",
            "K1: к пролонгации",
            "K1: пролонгировано",
            "K1",
            "K2: к пролонгации",
            "K2: пролонгировано",
            "K2",
        ),
        rows,
    )
    _autosize_columns(ws)


def _write_exclusions_sheet(wb: Workbook, result):
    rows = []
    for pid in sorted(result.exclude_stop_set):
        row = result.proj.loc[pid]
        rows.append(
            (
                pid,
                row.get("manager"),
                row.get("month"),
                result.exclusion_reasons.get(pid, "Исключён из базы расчёта"),
            )
        )

    ws = wb.create_sheet("Исключенные проекты")
    _write_rows(ws, ("id", "Менеджер", "Последний месяц", "Причина"), rows)
    _autosize_columns(ws)


def build_excel_report(result, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_dept = wb.active
    ws_dept.title = "Весь отдел"
    _write_primer_table(ws_dept, "Месяц", result.dept_rows, top=1)

    ws_year = wb.create_sheet("Менеджеры за год")
    _write_primer_table(ws_year, "Менеджер", result.mgr_rows, top=1)

    ws_month = wb.create_sheet("Менеджеры по месяцам")
    col_labels = list(result.mat_k1.columns)
    k2_top = 1 + 3 + len(result.managers) + 2
    _write_coef_block(ws_month, "Коэффициент 1", col_labels, result.managers, result.mat_k1, top=1)
    _write_coef_block(ws_month, "Коэффициент 2", col_labels, result.managers, result.mat_k2, top=k2_top)

    _write_summary_sheet(wb, result)
    _write_manager_ranking_sheet(wb, result)
    _write_exclusions_sheet(wb, result)

    for ws in wb.worksheets[:3]:
        _autosize_columns(ws)

    wb.save(output_path)
    return output_path
